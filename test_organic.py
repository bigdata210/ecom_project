import tempfile
import shutil
import unittest
from pathlib import Path
import openpyxl
from organic_parser import merge_pdfs,parse_organic
from organic_app import write_workbook

BASE=Path('C:/Users/hanyoung/OneDrive/Desktop/환경프로젝트')

class OrganicTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):cls.r=merge_pdfs([BASE/'VOC데이터.pdf',BASE/'유기인데이터.pdf'])
    def test_detection_and_split_digits(self):
        self.assertEqual(len(self.r),10)
        m={r['sheet']:r for r in self.r}
        self.assertEqual(m['m,p-xylene']['standards'][-1]['concentration'],200)
        self.assertEqual(m['m,p-xylene']['standards'][-1]['area'],1122986.625)
        self.assertEqual(m['Diazinon']['standards'][-1]['area'],1113.37146)
        self.assertTrue(all(s['area'] is None for r in self.r for s in r['samples']))
    def test_duplicate_rejected(self):
        with self.assertRaises(ValueError):merge_pdfs([BASE/'VOC데이터.pdf',BASE/'VOC데이터.pdf'])
    def test_identifies_content_with_arbitrary_filename(self):
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'report_001.pdf'
            shutil.copyfile(BASE/'유기인데이터.pdf',path)
            result=parse_organic(path)
            self.assertEqual([r['kind'] for r in result],['OP','OP'])
    def test_output_and_preservation(self):
        with tempfile.TemporaryDirectory() as d:
            out=Path(d)/'result.xlsx';write_workbook(self.r,'templates/organic_template.xlsx',out)
            w=openpyxl.load_workbook(out);values=openpyxl.load_workbook(out,data_only=True)
            self.assertEqual(len(w.sheetnames),26)
            self.assertTrue(all(len(s._charts)==1 for s in w))
            self.assertEqual(w['Benzene']['B5'].value,15751.688)
            self.assertEqual(w['m,p-xylene']['A9'].value,200)
            self.assertEqual(values['Benzene']['E30'].value,None)
            self.assertEqual(w['Benzene']['H30'].value,'측정값 없음')
            self.assertEqual(w['Malathion']['B31'].value,None)
            self.assertEqual(w['Malathion']['E31'].value,None)
            self.assertIn('/1000',w['Benzene']['E30'].value)
            with self.assertRaises(FileExistsError):write_workbook(self.r,'templates/organic_template.xlsx',out)

if __name__=='__main__':unittest.main()
