import tempfile
import shutil
import unittest
from pathlib import Path
import openpyxl
from measurement_app import detect,process_all

BASE=Path('C:/Users/hanyoung/OneDrive/Desktop/환경프로젝트')

class UnifiedTests(unittest.TestCase):
    def test_detection_does_not_use_names(self):
        with tempfile.TemporaryDirectory() as d:
            for filename,kind in [('중금속데이터.pdf','중금속'),('유기인데이터.pdf','유기인'),('VOC데이터.pdf','VOC')]:
                p=Path(d)/'random.pdf';shutil.copyfile(BASE/filename,p)
                self.assertEqual(detect(p),kind)

    def test_three_inputs_make_three_xlsx_and_correct_blank_subtraction(self):
        with tempfile.TemporaryDirectory() as d:
            folder,report=process_all([BASE/'중금속데이터.pdf',BASE/'VOC데이터.pdf',BASE/'유기인데이터.pdf'],d)
            self.assertEqual(report['errors'],[])
            self.assertEqual(len(report['outputs']),3)
            first=openpyxl.load_workbook(next(folder.glob('*_1차.xlsx')),data_only=True)
            second=openpyxl.load_workbook(next(folder.glob('*_2차.xlsx')),data_only=True)
            formulas=openpyxl.load_workbook(next(folder.glob('*_2차.xlsx')))
            self.assertEqual(first['Conc. in Calib Units']['H10'].value,-.00050)
            self.assertEqual(first['Conc. in Calib Units RSDs']['H13'].value,'>999.9%')
            self.assertEqual(second['w.As']['C31'].value,-.00050)
            self.assertEqual(second['w.As']['D31'].value,.00939)
            self.assertAlmostEqual(second['w.As']['E31'].value,-.00989)
            self.assertAlmostEqual(second['W.Cd']['E31'].value,-.00053)
            self.assertAlmostEqual(second['W.Pb']['E31'].value,-.00299)
            self.assertEqual(formulas['w.As']['E31'].value,'=C31-D31')
            self.assertIsNone(second['W.Zn']['E31'].value)
            self.assertIsNone(second['DATA']['C2'].value)
            organic=openpyxl.load_workbook(folder/'유기물_통합_01.xlsx',data_only=True)
            self.assertEqual(organic['Diazinon']['B5'].value,72.65617)
            self.assertEqual(organic['Benzene']['B5'].value,15751.688)

    def test_bad_file_does_not_block_valid_pdf(self):
        with tempfile.TemporaryDirectory() as d:
            invalid=Path(d)/'bad.pdf';invalid.write_text('not pdf')
            folder,report=process_all([invalid,BASE/'중금속데이터.pdf'],d)
            self.assertEqual(len(report['errors']),1)
            self.assertEqual(len(report['outputs']),2)

    def test_duplicate_compounds_make_separate_workbooks(self):
        with tempfile.TemporaryDirectory() as d:
            duplicate=Path(d)/'another_voc.pdf';shutil.copyfile(BASE/'VOC데이터.pdf',duplicate)
            folder,report=process_all([BASE/'VOC데이터.pdf',duplicate],d)
            self.assertEqual(len(report['outputs']),2)
            self.assertEqual(report['errors'],[])

if __name__=='__main__':unittest.main()
